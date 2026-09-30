from pathlib import Path

import pytest
from openpyxl import Workbook

from app.motores.compras_supply_chain.excel_local import ClienteExcelLocal
from app.motores.compras_supply_chain.google_sheets import (
    construir_fuente_compras_desde_entorno,
)


def _crear_excel(ruta: Path) -> None:
    workbook = Workbook()
    co = workbook.active
    co.title = "INFORME CLIENTES (CO)"
    encabezados = [
        "CLIENTE",
        "PROVEEDOR",
        "NUMERO OC",
        "ESTADO",
        "MODO TRANSPORTE",
        "VALOR TOTAL COMPRA USD",
        "VALOR OCI (DDP)",
    ]
    co.append(encabezados)
    co.append(["Cliente CO", "Proveedor", "OC-1", "EN PRODUCCION", "MARITIMO", 100, 120])

    for nombre, oc in (
        ("INFORME CLIENTES (EC)", "OC-2"),
        ("INFORME CLIENTES (CL)", "OC-3"),
    ):
        hoja = workbook.create_sheet(nombre)
        hoja.append(encabezados)
        hoja.append([nombre, "Proveedor", oc, "ENTREGADO", "AEREO", 50, 60])

    supply = workbook.create_sheet("Supply Chain ")
    supply.append(["ESTADO", "VALOR OCI (DDP)"])
    supply.append(["EN PRODUCCION", 120])
    workbook.save(ruta)


def test_cliente_excel_local_lee_hoja_con_espacios_y_rango_de_columnas(tmp_path: Path) -> None:
    ruta = tmp_path / "compras.xlsx"
    _crear_excel(ruta)
    cliente = ClienteExcelLocal(ruta)

    valores = cliente.obtener_valores(
        spreadsheet_id="ignorado",
        rango="'INFORME CLIENTES (CO)'!A:G",
    )

    assert valores[0][:5] == [
        "CLIENTE",
        "PROVEEDOR",
        "NUMERO OC",
        "ESTADO",
        "MODO TRANSPORTE",
    ]
    assert valores[1][2] == "OC-1"


def test_fuente_local_tiene_prioridad_sobre_demo(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    ruta = tmp_path / "compras.xlsx"
    _crear_excel(ruta)

    monkeypatch.setenv("APP_ENV", "development")
    monkeypatch.setenv("COMPRAS_DEMO_MODE", "true")
    monkeypatch.setenv("COMPRAS_EXCEL_LOCAL_FILE", str(ruta))
    monkeypatch.setenv("COMPRAS_SHEETS_EMPRESA_ID", "1")
    monkeypatch.setenv("COMPRAS_SHEETS_CO_RANGE", "'INFORME CLIENTES (CO)'!A:G")
    monkeypatch.setenv("COMPRAS_SHEETS_EC_RANGE", "'INFORME CLIENTES (EC)'!A:G")
    monkeypatch.setenv("COMPRAS_SHEETS_CL_RANGE", "'INFORME CLIENTES (CL)'!A:G")
    monkeypatch.setenv("COMPRAS_SHEETS_SUPPLY_CHAIN_RANGE", "'Supply Chain '!A:B")

    fuente = construir_fuente_compras_desde_entorno()
    snapshot = fuente.obtener_snapshot(empresa_id=1, forzar_lectura=True)

    assert fuente.configuracion.modo_fuente == "EXCEL_LOCAL"
    assert fuente.configuracion.spreadsheet_id == "excel-local:compras.xlsx"
    assert len(snapshot.lineas) == 3
    assert all(diagnostico.valido for diagnostico in snapshot.diagnosticos)
