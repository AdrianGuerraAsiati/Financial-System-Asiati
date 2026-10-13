"""Regresión de la Entrega 1: tablero, filtros, Excel y separación por empresa/moneda."""
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.core.usuarios.roles import ROL_ANALISTA_TESORERIA, ROL_TI
from tests.apoyo_auth import cliente_con_rol, cliente_superadmin, crear_empresa


def factura(cliente, empresa_id: int, numero: str, monto: str,
            vencimiento: date | None, *, moneda: str = "COP", proveedor: str = "Proveedor") -> dict:
    data = {
        "empresa_id": str(empresa_id), "numero": numero, "proveedor": proveedor,
        "monto": monto, "moneda": moneda, "descripcion": "Servicio " + numero,
    }
    if vencimiento is not None:
        data["fecha_vencimiento"] = vencimiento.isoformat()
    creado = cliente.post("/api/v1/tesoreria/facturas", data=data,
                          files={"archivo": (numero + ".pdf", b"%PDF-1.4\nfactura " + numero.encode(), "application/pdf")})
    assert creado.status_code == 201, creado.text
    return creado.json()


def abono(cliente, factura_id: int, monto: str, fecha: date, referencia: str):
    r = cliente.post(f"/api/v1/tesoreria/facturas/{factura_id}/pagos",
        data={"monto": monto, "fecha": fecha.isoformat(), "referencia": referencia},
        files={"comprobante": ("soporte-" + referencia + ".pdf",
                b"%PDF-1.4\n" + referencia.encode(), "application/pdf")})
    assert r.status_code == 201, r.text
    return r.json()


def datos_ejemplo(cliente, empresa_id: int):
    hoy = date.today()
    futuro = factura(cliente, empresa_id, "PAG-101", "1000", hoy + timedelta(days=12))
    vencida = factura(cliente, empresa_id, "VEN-202", "2000", hoy - timedelta(days=3))
    parcial = factura(cliente, empresa_id, "PAR-303", "3000", hoy + timedelta(days=2))
    completa = factura(cliente, empresa_id, "TOT-404", "500", hoy - timedelta(days=5))
    sin_fecha = factura(cliente, empresa_id, "PEN-505", "700", None)
    abono(cliente, parcial["id"], "1200", hoy, "PAR-AB1")
    abono(cliente, completa["id"], "500", hoy, "TOT-AB1")
    return [futuro, vencida, parcial, completa, sin_fecha]


def test_resumen_por_moneda_y_fecha_del_abono_no_factura(tmp_path, monkeypatch):
    monkeypatch.setenv("TESORERIA_DOCUMENTOS_DIR", str(tmp_path))
    cliente = cliente_superadmin()
    propia = crear_empresa("Tablero resumen")
    ajena = crear_empresa("Tablero otras empresas")
    datos_ejemplo(cliente, propia)
    factura(cliente, propia, "USD-600", "99.99", date.today(), moneda="USD")
    factura(cliente, ajena, "OTR-700", "900000", date.today())

    r = cliente.get("/api/v1/tesoreria/indicadores",
                    params={"empresa_id": propia, "moneda": "COP"})
    assert r.status_code == 200, r.text
    info = r.json()
    assert info["cantidad_facturas"] == 5
    assert info["total_pendiente"] == "5500.00"
    assert info["total_vencido"] == "2000.00"
    assert info["cantidad_vencidas"] == 1
    assert info["total_por_vencer_7d"] == "1800.00"
    assert info["cantidad_por_vencer_7d"] == 1
    assert info["sin_vencimiento"] == 1
    assert info["total_pagado_periodo"] == "1700.00"
    usd = cliente.get("/api/v1/tesoreria/indicadores",
                      params={"empresa_id": propia, "moneda": "USD"}).json()
    assert usd["total_pendiente"] == "99.99"
    assert usd["total_pagado_periodo"] == "0.00"

    pasado = (date.today() - timedelta(days=30)).isoformat()
    ayer = (date.today() - timedelta(days=1)).isoformat()
    periodo = cliente.get("/api/v1/tesoreria/indicadores",
                          params={"empresa_id": propia, "moneda": "COP",
                                  "pago_desde": pasado, "pago_hasta": ayer})
    assert periodo.status_code == 200
    assert periodo.json()["total_pagado_periodo"] == "0.00"


def test_filtros_paginacion_ambiguedad_y_validaciones(tmp_path, monkeypatch):
    monkeypatch.setenv("TESORERIA_DOCUMENTOS_DIR", str(tmp_path))
    cliente = cliente_superadmin()
    empresa = crear_empresa("Bandeja filtros")
    datos_ejemplo(cliente, empresa)

    base = {"empresa_id": empresa, "moneda": "COP"}
    pagina = cliente.get("/api/v1/tesoreria/consulta-facturas",
                        params={**base, "limite": 2, "desde": 0})
    assert pagina.status_code == 200, pagina.text
    assert pagina.json()["total"] == 5
    assert len(pagina.json()["items"]) == 2
    segunda = cliente.get("/api/v1/tesoreria/consulta-facturas",
                         params={**base, "limite": 2, "desde": 2})
    assert segunda.status_code == 200
    assert len({f["id"] for f in pagina.json()["items"] + segunda.json()["items"]}) == 4

    vencidas = cliente.get("/api/v1/tesoreria/consulta-facturas",
                         params={**base, "estado": "VENCIDA"}).json()
    assert vencidas["total"] == 1
    assert vencidas["items"][0]["numero"] == "VEN-202"

    parcial = cliente.get("/api/v1/tesoreria/consulta-facturas",
                         params={**base, "estado": "PARCIAL"}).json()
    assert parcial["total"] == 1
    assert parcial["items"][0]["alerta"] is None
    assert parcial["items"][0]["saldo"] == "1800.00"

    busqueda = cliente.get("/api/v1/tesoreria/consulta-facturas",
                          params={**base, "buscar": "ven-202"}).json()
    assert busqueda["total"] == 1

    fecha = (date.today() + timedelta(days=1)).isoformat()
    por_vencer = cliente.get("/api/v1/tesoreria/consulta-facturas",
                            params={**base, "vencimiento_desde": fecha})
    assert por_vencer.status_code == 200
    assert por_vencer.json()["total"] == 2

    invalido = cliente.get("/api/v1/tesoreria/consulta-facturas",
                          params={**base, "vencimiento_desde": "2026-12-01",
                                  "vencimiento_hasta": "2026-10-01"})
    assert invalido.status_code == 422
    assert cliente.get("/api/v1/tesoreria/consulta-facturas",
                      params={**base, "limite": 201}).status_code == 422


def test_excel_filtrado_muestra_facturas_y_pagos_sin_formulas(tmp_path, monkeypatch):
    monkeypatch.setenv("TESORERIA_DOCUMENTOS_DIR", str(tmp_path))
    cliente = cliente_superadmin()
    empresa = crear_empresa("Excel facturas")
    mala = factura(cliente, empresa, "EX-501", "1250.00", date.today(),
                   proveedor='=HYPERLINK("https://example.invalid","Click")')
    abono(cliente, mala["id"], "450", date.today(), "TRX-EX")
    factura(cliente, empresa, "USD-501", "75.00", date.today(), moneda="USD")

    resp = cliente.get("/api/v1/tesoreria/exportacion-facturas",
                       params={"empresa_id": empresa, "moneda": "COP", "buscar": "EX-501"})
    assert resp.status_code == 200, resp.text[:500]
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert resp.headers["cache-control"] == "no-store"
    libro = load_workbook(BytesIO(resp.content))
    assert libro.sheetnames == ["Facturas", "Pagos"]
    fc = libro["Facturas"]
    assert fc.max_row == 2
    assert fc["C2"].value == "EX-501"
    assert fc["D2"].data_type != "f"
    assert fc["D2"].value.startswith("'=HYPERLINK")
    assert Decimal(str(fc["H2"].value)) == Decimal("1250.00")
    assert Decimal(str(fc["I2"].value)) == Decimal("450.00")
    assert Decimal(str(fc["J2"].value)) == Decimal("800.00")
    pagos = libro["Pagos"]
    assert pagos.max_row == 2
    assert pagos["G2"].value == "TRX-EX"
    assert pagos["H2"].value == "soporte-TRX-EX.pdf"

    vacio = cliente.get("/api/v1/tesoreria/exportacion-facturas",
                       params={"empresa_id": empresa, "moneda": "CLP"})
    assert vacio.status_code == 200
    libro_vacio = load_workbook(BytesIO(vacio.content))
    assert libro_vacio["Facturas"].max_row == 1
    assert libro_vacio["Pagos"].max_row == 1


def test_endpoints_no_exponen_empresas_no_asignadas(tmp_path, monkeypatch):
    monkeypatch.setenv("TESORERIA_DOCUMENTOS_DIR", str(tmp_path))
    propia = crear_empresa("Empresa autorizada dashboard")
    ajena = crear_empresa("Empresa ajena dashboard")
    analista, _ = cliente_con_rol(ROL_ANALISTA_TESORERIA, empresas=(propia,))
    for ruta in ["indicadores", "consulta-facturas", "exportacion-facturas"]:
        assert analista.get("/api/v1/tesoreria/" + ruta,
                            params={"empresa_id": ajena}).status_code == 404
        assert analista.get("/api/v1/tesoreria/" + ruta,
                            params={"empresa_id": propia}).status_code == 200
    ti, _ = cliente_con_rol(ROL_TI)
    assert ti.get("/api/v1/tesoreria/indicadores", params={"empresa_id": ajena}).status_code == 200
