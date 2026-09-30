#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook, load_workbook


HOJAS = (
    ("INFORME CLIENTES (CO)", 59),
    ("INFORME CLIENTES (EC)", 59),
    ("INFORME CLIENTES (CL)", 59),
    ("Supply Chain ", 26),
)


def _fila_con_datos(valores: tuple[object, ...]) -> bool:
    return any(valor not in (None, "") for valor in valores)


def preparar(origen: Path, destino: Path) -> None:
    if not origen.is_file():
        raise SystemExit(f"No existe el archivo origen: {origen}")

    destino.parent.mkdir(parents=True, exist_ok=True)
    origen_wb = load_workbook(origen, read_only=True, data_only=True)
    destino_wb = Workbook(write_only=True)

    try:
        for nombre, max_columnas in HOJAS:
            if nombre not in origen_wb.sheetnames:
                raise SystemExit(f"Falta la hoja requerida {nombre!r} en {origen.name}.")

            ws_origen = origen_wb[nombre]
            ws_destino = destino_wb.create_sheet(nombre)

            ultima_fila_con_datos = 0
            filas = 0
            for numero_fila, fila in enumerate(
                ws_origen.iter_rows(
                    min_row=1,
                    min_col=1,
                    max_col=max_columnas,
                    values_only=True,
                ),
                start=1,
            ):
                if _fila_con_datos(fila):
                    ultima_fila_con_datos = numero_fila
                # Conservamos filas internas vacías para no desplazar números de fila.
                ws_destino.append(list(fila))
                filas += 1

                # En read_only algunos archivos reportan dimensiones exageradas por formato.
                # Tras 250 filas vacías consecutivas después de haber visto datos, cortamos.
                if (
                    ultima_fila_con_datos
                    and numero_fila - ultima_fila_con_datos >= 250
                ):
                    break

            print(
                f"{nombre}: {ultima_fila_con_datos} filas útiles "
                f"(se recorrieron {filas})"
            )
    finally:
        origen_wb.close()

    # Workbook(write_only=True) crea una hoja por cada create_sheet; no trae hoja inicial.
    destino_wb.save(destino)
    print(f"Excel operacional generado: {destino} ({destino.stat().st_size} bytes)")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    preparar(args.source, args.output)


if __name__ == "__main__":
    main()
