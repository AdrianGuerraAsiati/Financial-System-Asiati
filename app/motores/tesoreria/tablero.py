"""Consulta financiera de facturas: tablero, filtros paginados y Excel.

Todas las consultas se limitan a una empresa visible. Las monedas no se suman
entre sí; el histórico pagado usa fechas de abono, no la fecha de factura.
"""
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from io import BytesIO
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import case, func, or_, select
from sqlalchemy.orm import Session

from app.core.auth.dependencias import Acceso, requiere
from app.core.session import obtener_session
from app.motores.tesoreria.api import _empresa
from app.motores.tesoreria.model import FacturaCompra, PagoFactura
from app.motores.tesoreria.reglas import estado_factura

router = APIRouter(prefix="/tesoreria", tags=["tesoreria"])

Moneda = Literal["COP", "USD", "CLP"]
Estado = Literal["PAGADA", "PARCIAL", "PENDIENTE", "VENCIDA", "PENDIENTE_VALIDACION"]
MAX_EXPORTAR = 10000


@dataclass(frozen=True)
class Filtros:
    empresa_id: int
    moneda: Moneda
    buscar: str
    estado: Estado | None
    categoria: str | None
    vencimiento_desde: date | None
    vencimiento_hasta: date | None


def filtros_factura(
    empresa_id: int = Query(..., gt=0),
    moneda: Moneda = Query("COP"),
    buscar: str = Query("", max_length=120),
    estado: Estado | None = Query(None),
    categoria: str | None = Query(None, max_length=120),
    vencimiento_desde: date | None = Query(None),
    vencimiento_hasta: date | None = Query(None),
) -> Filtros:
    if vencimiento_desde and vencimiento_hasta and vencimiento_desde > vencimiento_hasta:
        raise HTTPException(422, "La fecha inicial no puede ser posterior a la final.")
    return Filtros(empresa_id, moneda, buscar.strip(), estado,
                   categoria or None, vencimiento_desde, vencimiento_hasta)


def _pagos_agrupados():
    return (
        select(
            PagoFactura.factura_id.label("factura_id"),
            func.sum(PagoFactura.monto).label("pagado"),
            func.max(PagoFactura.fecha).label("fecha_pago_real"),
            func.count(PagoFactura.id).label("cantidad_pagos"),
        ).group_by(PagoFactura.factura_id).subquery()
    )


def _consulta(f: Filtros, hoy: date):
    pagos = _pagos_agrupados()
    pagado = func.coalesce(pagos.c.pagado, Decimal("0.00"))
    saldo = FacturaCompra.monto - pagado
    estado = case(
        (saldo == 0, "PAGADA"),
        (pagado > 0, "PARCIAL"),
        (FacturaCompra.fecha_vencimiento.is_(None), "PENDIENTE_VALIDACION"),
        (FacturaCompra.fecha_vencimiento < hoy, "VENCIDA"),
        else_="PENDIENTE",
    )
    condiciones = [FacturaCompra.empresa_id == f.empresa_id, FacturaCompra.moneda == f.moneda]
    if f.buscar:
        literal = f.buscar.replace("!", "!!").replace("%", "!%").replace("_", "!_")
        patron = f"%{literal}%"
        condiciones.append(or_(
            FacturaCompra.numero.ilike(patron, escape="!"),
            FacturaCompra.proveedor.ilike(patron, escape="!"),
            FacturaCompra.descripcion.ilike(patron, escape="!"),
        ))
    if f.estado:
        condiciones.append(estado == f.estado)
    if f.categoria:
        condiciones.append(FacturaCompra.categoria == f.categoria)
    if f.vencimiento_desde:
        condiciones.append(FacturaCompra.fecha_vencimiento >= f.vencimiento_desde)
    if f.vencimiento_hasta:
        condiciones.append(FacturaCompra.fecha_vencimiento <= f.vencimiento_hasta)
    consulta = (
        select(FacturaCompra, pagado.label("pagado"),
               pagos.c.fecha_pago_real, pagos.c.cantidad_pagos)
        .outerjoin(pagos, FacturaCompra.id == pagos.c.factura_id)
        .where(*condiciones)
    )
    return consulta


def _item(factura: FacturaCompra, pagado: Decimal,
          fecha_pago_real: date | None, cantidad_pagos: int | None, hoy: date) -> dict:
    estado = estado_factura(factura.monto, [pagado], factura.fecha_vencimiento, hoy)
    return {
        "id": factura.id,
        "empresa_id": factura.empresa_id,
        "numero": factura.numero,
        "proveedor": factura.proveedor,
        "fecha_emision": factura.fecha_emision,
        "fecha_pago_oportuno": factura.fecha_vencimiento,
        "fecha_pago_real": fecha_pago_real,
        "monto": format(factura.monto, ".2f"),
        "moneda": factura.moneda,
        "descripcion": factura.descripcion,
        "ciudad": factura.ciudad,
        "unidad_negocio": factura.unidad_negocio,
        "categoria": factura.categoria,
        "fijo_variable": factura.fijo_variable,
        "cuenta_pago": factura.cuenta_pago,
        "documento_contable": factura.documento_contable,
        "pagos_cantidad": cantidad_pagos or 0,
        **estado,
    }


@router.get("/consulta-facturas")
def consulta_facturas(
    f: Filtros = Depends(filtros_factura),
    limite: int = Query(50, ge=1, le=200),
    desde: int = Query(0, ge=0),
    acceso: Acceso = Depends(requiere("tesoreria.ver")),
    session: Session = Depends(obtener_session),
) -> dict:
    _empresa(session, acceso, f.empresa_id)
    hoy = date.today()
    consulta = _consulta(f, hoy)
    total = session.scalar(select(func.count()).select_from(consulta.subquery())) or 0
    orden = (FacturaCompra.fecha_vencimiento.asc().nulls_last(),
             FacturaCompra.id.desc())
    filas = session.execute(consulta.order_by(*orden).offset(desde).limit(limite)).all()
    return {
        "items": [_item(factura, pagado, fecha, cantidad, hoy)
                  for factura, pagado, fecha, cantidad in filas],
        "total": total, "limite": limite, "desde": desde,
        "moneda": f.moneda,
    }


@router.get("/indicadores")
def indicadores(
    empresa_id: int = Query(..., gt=0),
    moneda: Moneda = Query("COP"),
    pago_desde: date | None = Query(None),
    pago_hasta: date | None = Query(None),
    acceso: Acceso = Depends(requiere("tesoreria.ver")),
    session: Session = Depends(obtener_session),
) -> dict:
    _empresa(session, acceso, empresa_id)
    hoy = date.today()
    inicio = pago_desde or hoy.replace(day=1)
    fin = pago_hasta or hoy
    if inicio > fin:
        raise HTTPException(422, "El rango de pagos es inválido.")
    pagos = _pagos_agrupados()
    pagado = func.coalesce(pagos.c.pagado, Decimal("0.00"))
    base = select(FacturaCompra.monto, FacturaCompra.fecha_vencimiento, pagado).outerjoin(
        pagos, FacturaCompra.id == pagos.c.factura_id
    ).where(FacturaCompra.empresa_id == empresa_id, FacturaCompra.moneda == moneda)
    pendiente = vencido = proximo = Decimal("0.00")
    num_vencidas = num_proximas = sin_vencimiento = num_facturas = 0
    # «Por vencer» es solo una ventana informativa de siete días calendario,
    # no una tolerancia ni un cambio de estado financiero.
    from datetime import timedelta
    hasta_proximo = hoy + timedelta(days=7)
    for monto, abonado_fecha, monto_pagado in session.execute(base):
        num_facturas += 1
        saldo = monto - monto_pagado
        if saldo <= 0:
            continue
        pendiente += saldo
        if abonado_fecha is None:
            sin_vencimiento += 1
        elif abonado_fecha < hoy:
            vencido += saldo
            num_vencidas += 1
        elif hoy <= abonado_fecha <= hasta_proximo:
            proximo += saldo
            num_proximas += 1

    total_pagado = session.scalar(select(func.coalesce(func.sum(PagoFactura.monto), 0))
        .join(FacturaCompra, FacturaCompra.id == PagoFactura.factura_id)
        .where(FacturaCompra.empresa_id == empresa_id,
               FacturaCompra.moneda == moneda,
               PagoFactura.fecha >= inicio, PagoFactura.fecha <= fin))
    return {
        "empresa_id": empresa_id, "moneda": moneda, "fecha_corte": hoy,
        "pago_desde": inicio, "pago_hasta": fin,
        "cantidad_facturas": num_facturas,
        "total_pendiente": format(pendiente, ".2f"),
        "total_vencido": format(vencido, ".2f"),
        "cantidad_vencidas": num_vencidas,
        "total_por_vencer_7d": format(proximo, ".2f"),
        "cantidad_por_vencer_7d": num_proximas,
        "sin_vencimiento": sin_vencimiento,
        "total_pagado_periodo": format(total_pagado, ".2f"),
    }


def _texto_excel(valor: object) -> str:
    # Evitar fórmulas maliciosas en valores de facturas/proveedores importados.
    texto = str(valor)[:32767] if valor is not None else ""
    if texto.lstrip().startswith(("=", "+", "-", "@")) or texto[:1] in ("\t", "\r", "\n"):
        return "'" + texto
    return texto


def _formato_hoja(hoja, columnas: list[tuple[str, int]], fechas: set[int],
                   montos: set[int]) -> None:
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    for indice, (titulo, ancho) in enumerate(columnas, start=1):
        celda = hoja.cell(1, indice, titulo)
        celda.fill = PatternFill("solid", fgColor="153B5C")
        celda.font = Font(bold=True, color="FFFFFF")
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    hoja.row_dimensions[1].height = 32
    for fila in hoja.iter_rows(min_row=2):
        for n in fechas:
            fila[n - 1].number_format = "dd/mm/yyyy"
        for n in montos:
            fila[n - 1].number_format = '#,##0.00'


@router.get("/exportacion-facturas")
def exportar_facturas(
    f: Filtros = Depends(filtros_factura),
    acceso: Acceso = Depends(requiere("tesoreria.ver")),
    session: Session = Depends(obtener_session),
):
    _empresa(session, acceso, f.empresa_id)
    hoy = date.today()
    consulta = _consulta(f, hoy)
    total = session.scalar(select(func.count()).select_from(consulta.subquery())) or 0
    if total > MAX_EXPORTAR:
        raise HTTPException(413, "Hay más de 10.000 facturas. Aplica filtros para exportar.")
    registros = session.execute(consulta.order_by(FacturaCompra.id)).all()

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Facturas"
    headers = [
        ("ID", 10), ("Empresa ID", 12), ("Número factura", 20), ("Proveedor", 34),
        ("Fecha emisión", 16), ("Fecha pago oportuno", 22), ("Fecha pago real", 20),
        ("Monto factura", 19), ("Pagado", 19), ("Saldo pendiente", 20),
        ("Moneda", 11), ("Descripción", 38), ("Ciudad", 19),
        ("Unidad negocio", 24), ("Categoría", 24), ("Cuenta pago", 28),
        ("Fijo / variable", 18), ("Doc. contable", 20),
        ("Estado", 24), ("Alerta", 24), ("Cantidad pagos", 17),
    ]
    hoja.append([h[0] for h in headers])
    facturas_ids = []
    for factura, pagado, fecha, cantidad in registros:
        estado = estado_factura(factura.monto, [pagado], factura.fecha_vencimiento, hoy)
        facturas_ids.append(factura.id)
        hoja.append([
            factura.id, factura.empresa_id, _texto_excel(factura.numero),
            _texto_excel(factura.proveedor), factura.fecha_emision,
            factura.fecha_vencimiento, fecha, factura.monto, pagado,
            factura.monto - pagado, _texto_excel(factura.moneda),
            _texto_excel(factura.descripcion), _texto_excel(factura.ciudad),
            _texto_excel(factura.unidad_negocio), _texto_excel(factura.categoria),
            _texto_excel(factura.cuenta_pago), _texto_excel(factura.fijo_variable),
            _texto_excel(factura.documento_contable), estado["estado"],
            estado["alerta"], cantidad or 0,
        ])
    _formato_hoja(hoja, headers, {5, 6, 7}, {8, 9, 10})
    pagos_hoja = libro.create_sheet("Pagos")
    cols_pagos = [
        ("Factura ID", 15), ("Número factura", 22), ("Proveedor", 34),
        ("Fecha pago real", 20), ("Monto pagado", 20), ("Moneda", 12),
        ("Referencia", 25), ("Comprobante", 35),
    ]
    pagos_hoja.append([h[0] for h in cols_pagos])
    if facturas_ids:
        pagos_registros = session.execute(
            select(PagoFactura, FacturaCompra.numero, FacturaCompra.proveedor, FacturaCompra.moneda)
            .join(FacturaCompra, FacturaCompra.id == PagoFactura.factura_id)
            .where(PagoFactura.factura_id.in_(facturas_ids))
            .order_by(PagoFactura.factura_id, PagoFactura.fecha, PagoFactura.id)
        )
        for pago, numero, proveedor, moneda in pagos_registros:
            pagos_hoja.append([
                pago.factura_id, _texto_excel(numero), _texto_excel(proveedor),
                pago.fecha, pago.monto, _texto_excel(moneda),
                _texto_excel(pago.referencia), _texto_excel(pago.soporte_nombre),
            ])
    _formato_hoja(pagos_hoja, cols_pagos, {4}, {5})
    archivo = BytesIO()
    libro.save(archivo)
    archivo.seek(0)
    nombre = f"Facturas_Tesoreria_empresa_{f.empresa_id}_{hoy:%Y%m%d}.xlsx"
    return StreamingResponse(
        archivo,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre}"',
                 "Cache-Control": "no-store"},
    )
