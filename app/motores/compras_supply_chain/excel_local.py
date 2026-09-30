from __future__ import annotations

from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, range_boundaries


_RANGO_CON_HOJA = re.compile(
    r"^(?:'((?:[^']|'')+)'|([^!]+))!(.+)$"
)


class LecturaExcelLocalError(RuntimeError):
    pass


def _separar_rango(rango: str) -> tuple[str, str]:
    match = _RANGO_CON_HOJA.match(rango.strip())
    if not match:
        raise LecturaExcelLocalError(
            f"Rango de Excel no soportado: {rango}. Usa Hoja!A:Z o 'Hoja con espacios'!A:Z."
        )
    hoja = (match.group(1) or match.group(2) or "").replace("''", "'")
    referencia = match.group(3).replace("$", "")
    return hoja, referencia


def _limites(referencia: str, max_row: int) -> tuple[int, int, int, int]:
    if ":" not in referencia:
        referencia = f"{referencia}:{referencia}"

    izquierda, derecha = referencia.split(":", 1)
    if izquierda.isalpha() and derecha.isalpha():
        return (
            column_index_from_string(izquierda),
            1,
            column_index_from_string(derecha),
            max_row,
        )

    min_col, min_row, max_col, max_row_ref = range_boundaries(referencia)
    return (
        min_col,
        min_row or 1,
        max_col,
        max_row_ref or max_row,
    )


class ClienteExcelLocal:
    """Lector read-only compatible con el contrato mínimo usado por Compras."""

    def __init__(self, archivo: str | Path) -> None:
        self.archivo = Path(archivo)

    def obtener_valores(
        self,
        *,
        spreadsheet_id: str,
        rango: str,
    ) -> list[list[Any]]:
        del spreadsheet_id
        if not self.archivo.is_file():
            raise LecturaExcelLocalError(
                f"No existe el Excel local de Compras: {self.archivo}."
            )

        hoja, referencia = _separar_rango(rango)
        try:
            workbook = load_workbook(
                self.archivo,
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            raise LecturaExcelLocalError(
                f"No se pudo abrir el Excel local de Compras: {self.archivo.name}."
            ) from exc

        try:
            if hoja not in workbook.sheetnames:
                raise LecturaExcelLocalError(
                    f"El Excel local no contiene la hoja {hoja!r}."
                )
            worksheet = workbook[hoja]
            min_col, min_row, max_col, max_row = _limites(
                referencia,
                worksheet.max_row,
            )
            resultado: list[list[Any]] = []
            for fila in worksheet.iter_rows(
                min_row=min_row,
                max_row=max_row,
                min_col=min_col,
                max_col=max_col,
                values_only=True,
            ):
                valores = list(fila)
                while valores and valores[-1] is None:
                    valores.pop()
                resultado.append(valores)

            while resultado and not any(
                valor not in (None, "") for valor in resultado[-1]
            ):
                resultado.pop()
            return resultado
        finally:
            workbook.close()

    def obtener_pivotes(
        self,
        *,
        spreadsheet_id: str,
        rango: str,
    ) -> tuple[()]:
        del spreadsheet_id, rango
        # openpyxl no expone de forma fiable la definición completa de pivotes.
        # La validación derivada puede operar sin metadata de pivote.
        return ()
